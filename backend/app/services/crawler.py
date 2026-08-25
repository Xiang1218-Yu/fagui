import hashlib
from collections import deque
from pathlib import Path
from typing import Optional
from urllib import robotparser
from urllib.parse import urldefrag, urljoin, urlparse

import httpx
from bs4 import BeautifulSoup
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import settings
from app.models import Attachment, Change, CrawlRun, Document, Snapshot, Source, utcnow
from app.services import dedup, diffutil, notify

ATTACH_EXTS = {".pdf", ".doc", ".docx", ".xls", ".xlsx"}


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _sha256_text(text: str) -> str:
    return _sha256_bytes(text.encode("utf-8"))


def _same_host(url: str, host: str) -> bool:
    return urlparse(url).netloc.lower() == host


def _path_allowed(url: str, prefixes: list) -> bool:
    path = urlparse(url).path or "/"
    return any(path.startswith(p) for p in prefixes)


def _load_robots(client: httpx.Client, base_url: str) -> Optional[robotparser.RobotFileParser]:
    """读取站点 robots.txt；获取失败视为全部允许（返回 None）。"""
    parsed = urlparse(base_url)
    robots_url = f"{parsed.scheme}://{parsed.netloc}/robots.txt"
    try:
        resp = client.get(robots_url)
        if resp.status_code != 200:
            return None
        rp = robotparser.RobotFileParser()
        rp.parse(resp.text.splitlines())
        return rp
    except Exception:
        return None


def extract_text(html: str):
    """提取标题与正文文本：去 script/style/nav/footer，优先 main/article，按行归一化空白。"""
    soup = BeautifulSoup(html, "lxml")
    title = soup.title.get_text(strip=True) if soup.title else ""
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    node = soup.find("main") or soup.find("article") or soup.body or soup
    lines = [ln.strip() for ln in node.get_text("\n").splitlines()]
    text = "\n".join(ln for ln in lines if ln)
    return title, text, soup


def _snapshot_dir() -> Path:
    day_dir = Path(settings.snapshot_dir) / utcnow().strftime("%Y-%m-%d")
    day_dir.mkdir(parents=True, exist_ok=True)
    return day_dir


def _store_snapshot(content_hash: str, raw_html: str, text: str):
    day_dir = _snapshot_dir()
    raw_path = day_dir / f"{content_hash}.html"
    text_path = day_dir / f"{content_hash}.txt"
    if not raw_path.exists():
        raw_path.write_text(raw_html, encoding="utf-8")
    if not text_path.exists():
        text_path.write_text(text, encoding="utf-8")
    return str(raw_path), str(text_path)


def _download(client: httpx.Client, url: str) -> Optional[bytes]:
    """流式下载附件，超过大小上限则放弃。"""
    chunks, total = [], 0
    with client.stream("GET", url) as resp:
        if resp.status_code >= 400:
            return None
        for chunk in resp.iter_bytes(65536):
            total += len(chunk)
            if total > settings.max_attachment_bytes:
                return None
            chunks.append(chunk)
    return b"".join(chunks)


def _extract_attachment_text(path: Path, ext: str) -> Optional[str]:
    """抽取附件文本：pdf 用 pypdf、docx 用 python-docx，失败返回 None。"""
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if ext == ".docx":
            import docx

            document = docx.Document(str(path))
            return "\n".join(p.text for p in document.paragraphs)
    except Exception:
        return None
    return None


def _read_text_file(path: Optional[str]) -> str:
    if path and Path(path).exists():
        try:
            return Path(path).read_text(encoding="utf-8")
        except OSError:
            return ""
    return ""


def _get_or_create_document(session: Session, source: Source, url: str, title: str):
    doc = session.scalar(select(Document).where(Document.source_id == source.id, Document.url == url))
    if doc is not None:
        if title and not doc.title:
            doc.title = title
        return doc, False
    doc = Document(source_id=source.id, url=url, title=title or url, created_at=utcnow())
    session.add(doc)
    session.flush()
    return doc, True


def _process_attachments(client: httpx.Client, soup: BeautifulSoup, page_url: str, doc: Document, run: CrawlRun, session: Session):
    """下载页面中的附件链接，返回 [(kind, attachment)]，kind 为 added/updated。"""
    updates = []
    seen = set()
    for a in soup.find_all("a", href=True):
        url = urldefrag(urljoin(page_url, a["href"]))[0]
        ext = Path(urlparse(url).path).suffix.lower()
        if ext not in ATTACH_EXTS or url in seen:
            continue
        seen.add(url)
        try:
            content = _download(client, url)
        except Exception:
            continue  # 单个附件失败不中断采集
        if not content:
            continue
        content_hash = _sha256_bytes(content)
        existing = session.scalar(select(Attachment).where(Attachment.document_id == doc.id, Attachment.url == url))
        if existing is not None and existing.content_hash == content_hash:
            continue  # 同 url 同 hash，跳过
        filename = Path(urlparse(url).path).name or f"attachment{ext}"
        day_dir = _snapshot_dir()
        snapshot_path = day_dir / f"att_{content_hash}{ext}"
        snapshot_path.write_bytes(content)
        text = _extract_attachment_text(snapshot_path, ext)
        text_path = None
        if text is not None:
            tp = day_dir / f"att_{content_hash}.txt"
            tp.write_text(text, encoding="utf-8")
            text_path = str(tp)
        kind = "updated" if existing is not None else "added"
        if existing is None:
            existing = Attachment(document_id=doc.id, url=url, created_at=utcnow())
            session.add(existing)
        existing.filename = filename
        existing.content_hash = content_hash
        existing.snapshot_path = str(snapshot_path)
        existing.text_path = text_path
        session.flush()
        run.attachments_fetched += 1
        updates.append((kind, existing))
    return updates


def crawl_source(source_id: int, session: Session) -> Optional[CrawlRun]:
    """采集单个来源的入口：并发保护 + 异常兜底。"""
    source = session.get(Source, source_id)
    if source is None:
        return None
    running = session.scalar(
        select(CrawlRun).where(CrawlRun.source_id == source_id, CrawlRun.status == "running")
    )
    if running is not None:
        return running  # 已有进行中的采集，直接返回
    run = CrawlRun(source_id=source_id, status="running", started_at=utcnow())
    session.add(run)
    session.commit()
    try:
        _crawl(source, run, session)
        run.status = "success"
    except Exception as exc:
        session.rollback()
        run.status = "failed"
        run.error = str(exc)[:1000]
    finally:
        run.finished_at = utcnow()
        session.commit()
    return run


def _crawl(source: Source, run: CrawlRun, session: Session) -> None:
    base_url = urldefrag(source.base_url)[0]
    host = urlparse(base_url).netloc.lower()
    prefixes = source.allowed_paths or ["/"]

    client = httpx.Client(
        headers={"User-Agent": settings.crawl_user_agent},
        timeout=15.0,
        follow_redirects=True,
    )
    try:
        rp = _load_robots(client, base_url) if source.respect_robots else None
        queue = deque([base_url])
        seen = {base_url}
        pages = 0
        while queue and pages < source.max_pages:
            url = queue.popleft()
            if not _same_host(url, host) or not _path_allowed(url, prefixes):
                continue
            if rp is not None and not rp.can_fetch(settings.crawl_user_agent, url):
                continue
            try:
                resp = client.get(url)
            except Exception:
                continue  # 单页失败不中断整体采集
            if resp.status_code >= 400:
                continue
            if "text/html" not in resp.headers.get("content-type", "").lower():
                continue
            pages += 1
            run.pages_fetched = pages

            html = resp.text
            title, text, soup = extract_text(html)
            doc, created = _get_or_create_document(session, source, url, title)
            old_snapshot = session.scalar(
                select(Snapshot).where(Snapshot.document_id == doc.id).order_by(Snapshot.id.desc()).limit(1)
            )

            new_snapshot = None
            content_hash = _sha256_text(text)
            if doc.latest_hash != content_hash:
                raw_path, text_path = _store_snapshot(content_hash, html, text)
                new_snapshot = Snapshot(
                    document_id=doc.id,
                    run_id=run.id,
                    raw_path=raw_path,
                    text_path=text_path,
                    content_hash=content_hash,
                    created_at=utcnow(),
                )
                session.add(new_snapshot)
                session.flush()
                doc.latest_hash = content_hash

            att_updates = _process_attachments(client, soup, url, doc, run, session)

            changes = []
            if created and new_snapshot is not None:
                changes.append(
                    Change(
                        document_id=doc.id,
                        change_type="new",
                        new_snapshot_id=new_snapshot.id,
                        diff_summary=f"首次收录：{title or url}",
                        detected_at=utcnow(),
                    )
                )
            else:
                if new_snapshot is not None:
                    old_text = _read_text_file(old_snapshot.text_path if old_snapshot else None)
                    changes.append(
                        Change(
                            document_id=doc.id,
                            change_type="content_updated",
                            old_snapshot_id=old_snapshot.id if old_snapshot else None,
                            new_snapshot_id=new_snapshot.id,
                            diff_summary=diffutil.diff_summary(old_text, text),
                            detected_at=utcnow(),
                        )
                    )
                for kind, att in att_updates:
                    label = "新增" if kind == "added" else "内容更新"
                    changes.append(
                        Change(
                            document_id=doc.id,
                            change_type="attachment_updated",
                            diff_summary=f"附件《{att.filename}》{label}",
                            detected_at=utcnow(),
                        )
                    )

            for change in changes:
                regulation = dedup.assign_regulation(session, doc, text)
                change.regulation_id = regulation.id
                session.add(change)
                session.flush()
                notify.dispatch_for_change(session, change, doc, regulation, source)
                run.changes_detected += 1

            # BFS 发现同站新链接
            for a in soup.find_all("a", href=True):
                next_url = urldefrag(urljoin(url, a["href"]))[0]
                if next_url in seen or not next_url.startswith(("http://", "https://")):
                    continue
                seen.add(next_url)
                if _same_host(next_url, host) and _path_allowed(next_url, prefixes):
                    queue.append(next_url)
    finally:
        client.close()
