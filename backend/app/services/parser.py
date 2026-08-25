import re
import hashlib
import logging
from datetime import datetime
from typing import Optional, Any
from bs4 import BeautifulSoup
from urllib.parse import urljoin, urlparse
from dateutil import parser as dateutil_parser

logger = logging.getLogger(__name__)


class ParsedContent:
    def __init__(
        self,
        title: str,
        text: str,
        html: str,
        links: list[dict[str, str]],
        attachments: list[dict[str, str]],
        metadata: dict[str, Any],
    ):
        self.title = title
        self.text = text
        self.html = html
        self.links = links
        self.attachments = attachments
        self.metadata = metadata
        self.content_hash = hashlib.sha256(text.encode("utf-8")).hexdigest()
        self.word_count = len(text.split())


class ContentParser:
    ATTACHMENT_EXTENSIONS = {
        ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx",
        ".txt", ".csv", ".rtf", ".zip", ".rar", ".7z",
    }

    def __init__(self, selector_config: Optional[dict[str, Any]] = None):
        self.selector_config = selector_config or {}

    def parse(self, html: str, base_url: str) -> ParsedContent:
        soup = BeautifulSoup(html, "lxml")

        for element in soup(["script", "style", "nav", "footer", "header", "noscript"]):
            element.decompose()

        title = self._extract_title(soup)
        content_html, content_text = self._extract_content(soup)
        links = self._extract_links(soup, base_url)
        attachments = self._extract_attachments(soup, base_url)
        metadata = self._extract_metadata(soup)

        return ParsedContent(
            title=title,
            text=content_text,
            html=content_html,
            links=links,
            attachments=attachments,
            metadata=metadata,
        )

    def _extract_title(self, soup: BeautifulSoup) -> str:
        title = ""
        title_selectors = self.selector_config.get("title", [
            "h1", ".article-title", ".content-title", ".title", "title"
        ])
        for selector in title_selectors:
            el = soup.select_one(selector)
            if el:
                title = el.get_text(strip=True)
                if title:
                    break
        if not title:
            if soup.title:
                title = soup.title.get_text(strip=True)
        return title[:1024]

    def _extract_content(self, soup: BeautifulSoup) -> tuple[str, str]:
        content_el = None
        content_selectors = self.selector_config.get("content", [
            "article", ".article-content", ".content-body", ".main-content",
            "#content", ".content", "main", ".post-content", ".entry-content"
        ])
        for selector in content_selectors:
            content_el = soup.select_one(selector)
            if content_el:
                break

        if not content_el:
            content_el = soup.body if soup.body else soup

        content_html = str(content_el)
        text = content_el.get_text(separator="\n", strip=True)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r"[ \t]+", " ", text)
        return content_html, text.strip()

    def _extract_links(self, soup: BeautifulSoup, base_url: str) -> list[dict[str, str]]:
        links = []
        for a in soup.find_all("a", href=True):
            href = a["href"].strip()
            if href.startswith("#") or href.startswith("javascript:") or href.startswith("mailto:"):
                continue
            absolute_url = urljoin(base_url, href)
            parsed = urlparse(absolute_url)
            if parsed.scheme in ("http", "https"):
                links.append({
                    "url": absolute_url,
                    "text": a.get_text(strip=True)[:512],
                })
        return links

    def _extract_attachments(self, soup: BeautifulSoup, base_url: str) -> list[dict[str, str]]:
        attachments = []
        seen_urls = set()

        for a in soup.find_all("a", href=True):
            href = a["href"].strip()
            absolute_url = urljoin(base_url, href)
            parsed_path = urlparse(absolute_url).path.lower()
            if any(parsed_path.endswith(ext) for ext in self.ATTACHMENT_EXTENSIONS):
                if absolute_url not in seen_urls:
                    seen_urls.add(absolute_url)
                    attachments.append({
                        "url": absolute_url,
                        "filename": a.get_text(strip=True) or urlparse(href).path.split("/")[-1],
                    })
        return attachments

    def _extract_metadata(self, soup: BeautifulSoup) -> dict[str, Any]:
        metadata: dict[str, Any] = {}

        for meta in soup.find_all("meta"):
            name = meta.get("name", "").lower()
            prop = meta.get("property", "").lower()
            content = meta.get("content", "")
            if name == "description" or prop == "og:description":
                metadata["description"] = content[:1024]
            elif name == "keywords":
                metadata["keywords"] = content[:512]
            elif name == "author":
                metadata["author"] = content[:256]
            elif prop == "article:published_time":
                metadata["published_time"] = content

        for time_el in soup.find_all("time"):
            dt_str = time_el.get("datetime", "")
            if dt_str:
                try:
                    metadata["publish_date"] = dateutil_parser.parse(dt_str).isoformat()
                    break
                except (ValueError, TypeError):
                    pass

        return metadata

    def parse_attachment_text(self, content: bytes, content_type: str, filename: str) -> tuple[str, str]:
        ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
        try:
            if ext == "pdf" or "pdf" in content_type:
                return self._parse_pdf(content)
            elif ext in ("docx",) or "word" in content_type:
                return self._parse_docx(content)
            elif ext in ("xlsx",) or "excel" in content_type or "spreadsheet" in content_type:
                return self._parse_xlsx(content)
            elif ext in ("txt", "csv") or "text" in content_type:
                for encoding in ["utf-8", "gbk", "gb2312", "latin-1"]:
                    try:
                        return content.decode(encoding), "success"
                    except UnicodeDecodeError:
                        continue
                return content.decode("utf-8", errors="replace"), "partial"
            else:
                return "", f"unsupported_format: {ext}"
        except Exception as e:
            logger.error(f"Failed to parse attachment {filename}: {e}")
            return "", f"error: {str(e)[:200]}"

    def _parse_pdf(self, content: bytes) -> tuple[str, str]:
        try:
            import pdfplumber
            import io
            text_parts = []
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
            return "\n\n".join(text_parts), "success" if text_parts else "empty"
        except ImportError:
            return "", "pdfplumber_not_installed"
        except Exception as e:
            return "", f"pdf_error: {str(e)[:200]}"

    def _parse_docx(self, content: bytes) -> tuple[str, str]:
        try:
            from docx import Document
            import io
            doc = Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs), "success" if paragraphs else "empty"
        except ImportError:
            return "", "python-docx_not_installed"
        except Exception as e:
            return "", f"docx_error: {str(e)[:200]}"

    def _parse_xlsx(self, content: bytes) -> tuple[str, str]:
        try:
            from openpyxl import load_workbook
            import io
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            text_parts = []
            for sheet in wb.worksheets:
                text_parts.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join(str(c) for c in row if c is not None)
                    if row_text.strip():
                        text_parts.append(row_text)
            wb.close()
            return "\n".join(text_parts), "success" if text_parts else "empty"
        except ImportError:
            return "", "openpyxl_not_installed"
        except Exception as e:
            return "", f"xlsx_error: {str(e)[:200]}"
